"""Read the documents suppliers attach: PDF and XLSX to text, a quote PDF to structured terms."""

import io
import re

LABELS = {
    "Unit price (EUR)": "unit_price",
    "Payment terms (days net)": "payment_days",
    "Lead time (days)": "lead_time_days",
    "Volume per year (pcs)": "volume_per_year",
    "Contract term (months)": "contract_months",
    "Price valid from": "effective_date",
    "Index clause": "index_clause",
    "Incoterm": "incoterm",
    "Tooling (EUR)": "tooling_eur",
    "Quality commitment (ppm)": "ppm_commitment",
    "Chargeback per complaint (EUR)": "chargeback_eur_per_complaint",
}
SECTIONS = {
    "Offered terms": "stated",
    "Our standard terms": "standard",
    "Terms fixed by your enquiry": "fixed",
    "Agreed terms": "agreed",
}
CLAUSE = re.compile(r"(\S+), (\d+(?:[.,]\d+)?)% material, deadband ([\d.,]+) %, base (\S+)")
_CURRENCY = re.compile(r"\b(EUR|USD|GBP|CHF|CAD|AUD)\b|([€$£])", re.IGNORECASE)
_UNIT = re.compile(r"(?:/|per\s+)(\d+\s*)?(pcs?|pieces?|units?|kg|g|tonnes?|tons?|m)\b", re.IGNORECASE)


def read_document(name, data):
    """The text of an attachment: PDF pages, XLSX rows (formulas as written), or plain text."""
    if name.lower().endswith(".pdf"):
        from pypdf import PdfReader

        pages = [page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages]
        if pages and not any(page.strip() for page in pages):
            try:
                import pytesseract
                from pdf2image import convert_from_bytes
            except ImportError as exc:
                raise ValueError(
                    "PDF has pages without a text layer; install the optional OCR dependencies "
                    "and the Tesseract/Poppler system tools to read scanned pages"
                ) from exc
            try:
                images = convert_from_bytes(data)
                if len(images) != len(pages):
                    raise ValueError("OCR could not match rendered PDF pages to extracted pages")
                pages = [page.strip() or pytesseract.image_to_string(image)
                         for page, image in zip(pages, images)]
            except Exception as exc:
                if isinstance(exc, ValueError):
                    raise
                raise ValueError(f"OCR failed; check Tesseract and Poppler installation: {exc}") from exc
        text = "\n".join(pages)
        if not text.strip():
            raise ValueError("PDF contains no readable text, even after OCR")
        return text
    if name.lower().endswith(".xlsx"):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data))
        return "\n".join(
            f"[{ws.title}]\n"
            + "\n".join(" | ".join("" if v is None else str(v) for v in row) for row in ws.iter_rows(values_only=True))
            for ws in wb.worksheets
        )
    return data.decode("utf-8", "replace")


def _value(field, text):
    if field == "index_clause":
        m = CLAUSE.match(text)
        if not m:
            return None
        series, share, deadband, base = m.groups()
        return {
            "series": series,
            "material_share": float(share.replace(",", ".")) / 100,
            "deadband_pct": float(deadband.replace(",", ".")),
            "base_month": base,
        }
    if field in ("effective_date", "incoterm"):
        return text
    number = text.replace(" ", "").replace(",", ".")
    return float(number) if "eur" in field or field == "unit_price" else int(float(number))


def quote_terms(text):
    """Terms per section of a quote or order confirmation PDF: stated, standard, fixed, agreed."""
    out, section = {}, None
    for line in text.splitlines():
        line = line.strip()
        section = next((key for head, key in SECTIONS.items()
                        if line.casefold().startswith(head.casefold())), section)
        matched = next((match for lab in sorted(LABELS, key=len, reverse=True)
                        if (match := _match_label(line, lab))), None)
        if section and matched:
            field, label_unit, raw_value = matched
            values = out.setdefault(section, {})
            values[field] = _value(field, raw_value)
            if field == "unit_price":
                values["currency"] = _currency(label_unit) or values.get("currency")
                values["price_unit"] = _price_unit(label_unit) or "each"
    return out


def _match_label(line, label):
    base = label.split(" (", 1)[0]
    match = re.match(rf"^\s*{re.escape(base)}(?:\s*\(([^)]*)\))?\s*:?\s*(.*?)\s*$",
                     line, re.IGNORECASE)
    if not match:
        return None
    field = LABELS[label]
    return field, match.group(1) or "", match.group(2)


def _currency(label_unit):
    match = _CURRENCY.search(label_unit or "")
    if not match:
        return None
    code, symbol = match.groups()
    return code.upper() if code else {"€": "EUR", "$": "USD", "£": "GBP"}.get(symbol)


def _price_unit(label_unit):
    match = _UNIT.search(label_unit or "")
    if not match:
        return None
    quantity, unit = match.groups()
    unit = unit.casefold().rstrip("s")
    if unit in {"pc", "piece", "unit"}:
        unit = "each"
    return f"{quantity.strip()} {unit}" if quantity else unit


def standard_terms(events):
    """The supplier's standard terms from the latest quote PDF in a thread (empty if none)."""
    for msg in reversed(events):
        for att in msg.get("attachments") or []:
            if att.get("name", "").startswith("quote") and att.get("text"):
                return quote_terms(att["text"]).get("standard", {})
    return {}
